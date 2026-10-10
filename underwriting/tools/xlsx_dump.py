#!/usr/bin/env python3
"""Dump every sheet of an .xlsx workbook as JSON {sheet name: [[cell, ...], ...]}.

Used by import_grid.js so the condition grid can be imported straight from
Excel, all tabs at once. Merged cells are filled with the merged value so
headers that span columns still label every column.
"""
import json
import sys

import openpyxl


def sheet_rows(ws):
    merged = {}
    for rng in ws.merged_cells.ranges:
        value = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = value
    rows = []
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        out = []
        for c, value in enumerate(row, start=1):
            if value is None:
                value = merged.get((r, c))
            out.append("" if value is None else str(value).strip())
        rows.append(out)
    return rows


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    print(json.dumps({ws.title: sheet_rows(ws) for ws in wb.worksheets}))


if __name__ == "__main__":
    main(sys.argv[1])
